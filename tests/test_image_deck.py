"""Contract tests for image-only delivery, without requiring PowerPoint."""
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "images_to_pptx.py"


class ImageDeckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.images = self.root / "slides"
        self.images.mkdir()

    def make_image(self, name, color, size=(320, 180)):
        path = self.images / name
        Image.new("RGB", size, color).save(path)
        return path

    def run_builder(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)],
                              capture_output=True, text=True)

    def test_natural_order_one_original_image_per_slide(self):
        paths = [self.make_image(name, color) for name, color in
                 [("10.png", "blue"), ("2.jpg", "green"), ("1.png", "red")]]
        output = self.root / "demo.pptx"
        result = self.run_builder(self.images, "-o", output)
        self.assertEqual(result.returncode, 0, result.stderr)
        deck = Presentation(output)
        self.assertEqual(len(deck.slides), 3)
        for slide, path in zip(deck.slides, [paths[2], paths[1], paths[0]]):
            self.assertEqual(len(slide.shapes), 1)
            picture = slide.shapes[0]
            self.assertEqual(picture.shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertEqual(picture.image.blob, path.read_bytes())
            self.assertEqual((picture.left, picture.top), (0, 0))
            self.assertEqual((picture.width, picture.height),
                             (deck.slide_width, deck.slide_height))
            self.assertEqual((picture.crop_left, picture.crop_right,
                              picture.crop_top, picture.crop_bottom), (0, 0, 0, 0))

    def test_existing_delivery_is_not_overwritten(self):
        self.make_image("01.png", "red")
        output = self.root / "demo.pptx"
        output.write_bytes(b"existing delivery")
        result = self.run_builder(self.images, "-o", output)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("exists", result.stderr.lower())
        self.assertEqual(output.read_bytes(), b"existing delivery")

    def test_mixed_aspect_ratios_are_rejected_before_delivery(self):
        self.make_image("01.png", "red")
        self.make_image("02.png", "blue", (180, 320))
        output = self.root / "demo.pptx"
        result = self.run_builder(self.images, "-o", output)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("aspect ratio", result.stderr.lower())
        self.assertFalse(output.exists())

    def test_explicit_file_order_is_respected(self):
        a = self.make_image("a.png", "red")
        b = self.make_image("b.png", "blue")
        output = self.root / "demo.pptx"
        result = self.run_builder(b, a, "-o", output)
        self.assertEqual(result.returncode, 0, result.stderr)
        deck = Presentation(output)
        self.assertEqual(deck.slides[0].shapes[0].image.blob, b.read_bytes())
        self.assertEqual(deck.slides[1].shapes[0].image.blob, a.read_bytes())

    def test_broken_image_does_not_leave_partial_pptx(self):
        self.make_image("01.png", "red")
        (self.images / "02.png").write_bytes(b"not a PNG")
        output = self.root / "demo.pptx"
        result = self.run_builder(self.images, "-o", output)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("invalid image", result.stderr.lower())
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()

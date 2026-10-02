#!/usr/bin/env python3
"""PPT workflow 3.0: one original image per slide, with no reconstruction."""
from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path
import re
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}


def natural_key(path: Path):
    return [int(part) if part.isdigit() else part.casefold()
            for part in re.split(r"(\d+)", path.name)]


def collect_images(inputs: list[Path]) -> list[Path]:
    if len(inputs) == 1 and inputs[0].is_dir():
        images = sorted((p for p in inputs[0].iterdir()
                         if p.is_file() and not p.name.startswith(".")
                         and p.suffix.lower() in IMAGE_SUFFIXES), key=natural_key)
    else:
        images = inputs
    if not images:
        raise ValueError("No PNG/JPEG images found.")
    for image in images:
        if not image.is_file() or image.suffix.lower() not in IMAGE_SUFFIXES:
            raise ValueError(f"Not a PNG/JPEG image file: {image}")
    return images


def read_image(path: Path):
    blob = path.read_bytes()
    try:
        with Image.open(BytesIO(blob)) as image:
            image.load()
            if image.format not in {"PNG", "JPEG"}:
                raise ValueError("only PNG and JPEG are supported")
            if getattr(image, "n_frames", 1) != 1:
                raise ValueError("animated images are not supported")
            if image.getexif().get(274, 1) != 1:
                raise ValueError("export the image without EXIF rotation first")
            width, height = image.size
    except Exception as error:
        raise ValueError(f"Invalid image {path.name}: {error}") from error
    return blob, width, height


def make_deck(images: list[Path]) -> bytes:
    source = [read_image(path) for path in images]
    first_ratio = source[0][1] / source[0][2]
    for path, (_, width, height) in zip(images, source):
        if abs((width / height) / first_ratio - 1) > 0.001:
            raise ValueError(f"Different aspect ratio: {path.name}. "
                             "Use full-slide images with the same aspect ratio.")
    deck = Presentation()
    deck.slide_width = Inches(13 + 1 / 3)
    deck.slide_height = round(deck.slide_width / first_ratio)
    deck.core_properties.title = "PPT Workflow 3.0 | Image Only"
    deck.core_properties.subject = "One original image per slide"
    deck.core_properties.author = ""
    deck._element.set("autoCompressPictures", "0")
    for number, (path, (blob, width, height)) in enumerate(zip(images, source), 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        # Preserve aspect ratio even for a rounding difference of one source pixel.
        scale = min(deck.slide_width / width, deck.slide_height / height)
        picture_width, picture_height = round(width * scale), round(height * scale)
        picture = slide.shapes.add_picture(
            BytesIO(blob),
            (deck.slide_width - picture_width) // 2,
            (deck.slide_height - picture_height) // 2,
            width=picture_width,
            height=picture_height,
        )
        picture.name = f"Slide {number:02d} | {path.name}"
    buffer = BytesIO()
    deck.save(buffer)
    result = buffer.getvalue()
    verify_deck(result, [entry[0] for entry in source])
    return result


def verify_deck(blob: bytes, originals: list[bytes]):
    deck = Presentation(BytesIO(blob))
    if len(deck.slides) != len(originals):
        raise RuntimeError("Verification failed: unexpected slide count.")
    for number, (slide, original) in enumerate(zip(deck.slides, originals), 1):
        if len(slide.shapes) != 1:
            raise RuntimeError(f"Verification failed: slide {number} is not one image.")
        picture = slide.shapes[0]
        if picture.image.blob != original:
            raise RuntimeError(f"Verification failed: image {number} was changed.")
        if any((picture.crop_left, picture.crop_right,
                picture.crop_top, picture.crop_bottom)):
            raise RuntimeError(f"Verification failed: image {number} was cropped.")


def write_output(output: Path, blob: bytes, overwrite: bool):
    if output.exists() and not overwrite:
        raise ValueError(f"Output exists: {output}. Choose a new name or use --force.")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".pptx", delete=False) as file:
            temporary = Path(file.name)
            file.write(blob)
        if output.exists() and not overwrite:
            raise ValueError(f"Output exists: {output}.")
        os.replace(temporary, output)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("images", nargs="+", type=Path,
                        help="One image folder (natural filename order), or files in explicit order")
    parser.add_argument("-o", "--output", type=Path, required=True, help="Output .pptx file")
    parser.add_argument("--force", action="store_true", help="Replace an existing output")
    args = parser.parse_args()
    try:
        if args.output.suffix.lower() != ".pptx":
            raise ValueError("Output filename must end with .pptx.")
        if args.output.exists() and not args.force:
            raise ValueError(f"Output exists: {args.output}. Choose a new name or use --force.")
        images = collect_images(args.images)
        blob = make_deck(images)
        write_output(args.output, blob, args.force)
    except (ValueError, OSError, RuntimeError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(f"Created {args.output} ({len(images)} slides).")
    print("Verified: one image per slide; embedded image bytes match every source.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
